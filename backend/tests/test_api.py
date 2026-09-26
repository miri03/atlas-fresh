"""API tests: seed/plan endpoints, validation mapping and SPA serving."""

from __future__ import annotations

import os
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

os.environ.setdefault("ATLAS_DATA_PATH", str(Path(__file__).resolve().parents[2] / "data" / "Atlas_Fresh_Production_Commercial_Data.xlsx"))

from app.main import app  # noqa: E402

client = TestClient(app)


def test_health():
    resp = client.get("/api/health")
    assert resp.status_code == 200
    assert resp.json()["status"] == "ok"


def test_seed_returns_health():
    resp = client.post("/api/seed")
    assert resp.status_code == 200
    body = resp.json()
    assert body["actual_total_t"] == 560.0
    assert body["expected_total_t"] == 600.0
    assert body["station_capacity_t"] == 500.0
    assert body["actual_by_segment_t"] == {"A": 90.0, "B": 160.0, "C": 180.0, "D": 130.0}


def test_plan_matches_public_baseline():
    resp = client.post("/api/plan")
    assert resp.status_code == 200
    body = resp.json()
    k = body["kpis"]
    assert k["expected_plan_t"] == 600.0
    assert k["actual_received_t"] == 560.0
    assert k["export_t"] == 500.0
    assert k["local_t"] == 60.0
    assert k["export_rate_pct"] == 89.3
    assert k["export_revenue_eur"] == 549500.0
    assert k["local_value_eur"] == 4500.0
    assert k["total_value_eur"] == 554000.0
    assert k["at_risk_client_count"] == 3


def test_every_allocation_traceable():
    resp = client.post("/api/plan")
    body = resp.json()
    known_farms = {a["farm_id"] for a in body["allocations"]}
    known_clients = {a["client_id"] for a in body["allocations"]}
    assert all(f.startswith("F") for f in known_farms)
    assert all(c.startswith("C") for c in known_clients)
    assert all(a["segment"] in ("A", "B", "C", "D") for a in body["allocations"])


def test_plan_rejects_invalid_workbook_via_api(monkeypatch, tmp_path):
    import shutil

    from app.loader import WorkbookValidationError, load_workbook

    bad = tmp_path / "bad.xlsx"
    shutil.copy(Path(os.environ["ATLAS_DATA_PATH"]), bad)

    import openpyxl

    wb = openpyxl.load_workbook(bad)
    ws = wb["Clients"]
    for i, row in enumerate(ws.iter_rows(), start=1):
        if row[0].value == "C01":
            ws.cell(row=i, column=5).value = 50.5  # not a multiple of 5
            break
    wb.save(bad)

    original = load_workbook
    with monkeypatch.context() as m:
        m.setenv("ATLAS_DATA_PATH", str(bad))

        def fake_load(path):
            raise WorkbookValidationError([("Clients:C01", "demand_t=50.5 is not a multiple of 5 t")])

        m.setattr("app.main.load_workbook", fake_load)
        m.setattr("app.main.DEFAULT_DATA_PATH", bad)
        resp = client.post("/api/plan")
        assert resp.status_code == 422
        body = resp.json()
        assert body["error"] == "workbook_invalid"
        assert body["issues"][0]["location"] == "Clients:C01"