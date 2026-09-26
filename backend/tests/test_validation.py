"""Workbook validation tests: the loader rejects bad data loudly."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest

from app.loader import WorkbookValidationError, load_workbook


@pytest.fixture
def copy_workbook(workbook_path, tmp_path) -> Path:
    dest = tmp_path / "test.xlsx"
    shutil.copy(workbook_path, dest)
    return dest


def _rewrite_cell(path: Path, sheet: str, row: int, col: int, value) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    ws.cell(row=row, column=col).value = value
    wb.save(path)


def _find_row(path: Path, sheet: str, col: int, needle: str) -> int:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row[col - 1] is not None and str(row[col - 1]).strip() == needle:
            return i
    raise AssertionError(f"{needle} not found in {sheet}")


def _issues(e: WorkbookValidationError) -> list[tuple[str, str]]:
    return e.issues


def test_valid_workbook_passes(workbook_path):
    data = load_workbook(workbook_path)
    assert len(data.farms) == 20
    assert len(data.clients) == 10
    assert data.station.export_conditioning_capacity_t == 500


def test_duplicate_farm_id_rejected(copy_workbook):
    row = _find_row(copy_workbook, "Farms", 1, "F02")
    _rewrite_cell(copy_workbook, "Farms", row, 1, "F01")
    with pytest.raises(WorkbookValidationError) as exc:
        load_workbook(copy_workbook)
    assert any("duplicate farm_id" in msg and "F01" in loc for loc, msg in _issues(exc.value))


def test_duplicate_client_id_rejected(copy_workbook):
    row = _find_row(copy_workbook, "Clients", 1, "C02")
    _rewrite_cell(copy_workbook, "Clients", row, 1, "C01")
    with pytest.raises(WorkbookValidationError) as exc:
        load_workbook(copy_workbook)
    assert any("duplicate client_id" in msg and "C01" in loc for loc, msg in _issues(exc.value))


def test_mix_not_summing_to_one_rejected(copy_workbook):
    row = _find_row(copy_workbook, "Farms", 1, "F01")
    _rewrite_cell(copy_workbook, "Farms", row, 4, 0.2)  # expected_B_pct -> mix sums to 1.1
    with pytest.raises(WorkbookValidationError) as exc:
        load_workbook(copy_workbook)
    assert any("sum" in msg.lower() and "F01" in loc for loc, msg in _issues(exc.value))


def test_mix_fraction_out_of_range_rejected(copy_workbook):
    row = _find_row(copy_workbook, "Farms", 1, "F01")
    _rewrite_cell(copy_workbook, "Farms", row, 4, 1.5)
    with pytest.raises(WorkbookValidationError) as exc:
        load_workbook(copy_workbook)
    assert any("outside [0,1]" in msg and "F01" in loc for loc, msg in _issues(exc.value))


def test_negative_actual_rejected(copy_workbook):
    row = _find_row(copy_workbook, "Farms", 1, "F01")
    _rewrite_cell(copy_workbook, "Farms", row, 8, -5)
    with pytest.raises(WorkbookValidationError) as exc:
        load_workbook(copy_workbook)
    assert any("non-negative" in msg and "F01" in loc for loc, msg in _issues(exc.value))


def test_non_multiple_of_5_actual_rejected(copy_workbook):
    row = _find_row(copy_workbook, "Farms", 1, "F01")
    _rewrite_cell(copy_workbook, "Farms", row, 8, 27)
    with pytest.raises(WorkbookValidationError) as exc:
        load_workbook(copy_workbook)
    assert any("multiple of 5" in msg and "F01" in loc for loc, msg in _issues(exc.value))


def test_non_multiple_of_5_demand_rejected(copy_workbook):
    row = _find_row(copy_workbook, "Clients", 1, "C01")
    _rewrite_cell(copy_workbook, "Clients", row, 5, 52)
    with pytest.raises(WorkbookValidationError) as exc:
        load_workbook(copy_workbook)
    assert any("multiple of 5" in msg and "C01" in loc for loc, msg in _issues(exc.value))


def test_invalid_acceptance_mode_rejected(copy_workbook):
    row = _find_row(copy_workbook, "Clients", 1, "C01")
    _rewrite_cell(copy_workbook, "Clients", row, 3, "LOOSE")
    with pytest.raises(WorkbookValidationError) as exc:
        load_workbook(copy_workbook)
    assert any("acceptance_mode" in msg and "C01" in loc for loc, msg in _issues(exc.value))


def test_invalid_segment_rejected(copy_workbook):
    row = _find_row(copy_workbook, "Clients", 1, "C01")
    _rewrite_cell(copy_workbook, "Clients", row, 4, "E")
    with pytest.raises(WorkbookValidationError) as exc:
        load_workbook(copy_workbook)
    assert any("requested_segment" in msg and "C01" in loc for loc, msg in _issues(exc.value))


def test_invalid_capacity_rejected(copy_workbook):
    row = _find_row(copy_workbook, "Station", 1, "STATION-01")
    _rewrite_cell(copy_workbook, "Station", row, 2, -100)
    with pytest.raises(WorkbookValidationError) as exc:
        load_workbook(copy_workbook)
    assert any("capacity" in msg.lower() and "STATION-01" in loc for loc, msg in _issues(exc.value))


def test_missing_reference_price_rejected(copy_workbook):
    row = _find_row(copy_workbook, "Station", 1, "C")
    _rewrite_cell(copy_workbook, "Station", row, 2, None)
    with pytest.raises(WorkbookValidationError) as exc:
        load_workbook(copy_workbook)
    assert any("missing reference export price" in msg and "segment C" in msg for _, msg in _issues(exc.value))


def test_negative_demand_rejected(copy_workbook):
    row = _find_row(copy_workbook, "Clients", 1, "C01")
    _rewrite_cell(copy_workbook, "Clients", row, 5, -5)
    with pytest.raises(WorkbookValidationError) as exc:
        load_workbook(copy_workbook)
    assert any("non-negative" in msg and "C01" in loc for loc, msg in _issues(exc.value))